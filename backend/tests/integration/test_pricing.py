"""M5-09: the pricing step writes an XLSX built only from the tenant's rate card, stores
it and records a pricing_template artifact; POST /pursuits/{id}/agents/run drives the
pipeline and pauses at the first step a later task still has to add."""

from __future__ import annotations

import io
import uuid
from decimal import Decimal
from pathlib import Path
from typing import Any

import httpx
import openpyxl
import pytest
from app.agents import pipeline
from app.agents.pricing import XLSX_CONTENT_TYPE, PricingOutput, pricing_template_key
from app.agents.runner import AgentRunner
from app.agents.services import AgentServices
from app.core.compliance import ARTIFACT_PRICING_TEMPLATE
from app.core.config import Region, Settings
from app.core.db import Database
from app.core.opportunity import NoticeType
from app.core.pricing import NEEDS_CATEGORY, NEEDS_PRICE
from app.core.profile_fields import RateUnit
from app.core.roles import Role
from app.models import AgentStep, PursuitArtifact, RateCardEntry
from app.services.scanner import NoopScanner
from app.services.storage import LocalStorage, StorageRouter
from sqlalchemy import select

from tests.auth import auth_headers
from tests.integration.test_matrix import US_REQUIREMENTS, _setup
from tests.llm_fake import FakeLLM

SETTINGS = Settings(_env_file=None)  # type: ignore[call-arg]


def _services(tmp_path: Path, region: Region = Region.US) -> AgentServices:
    storage = LocalStorage(tmp_path, f"bidradar-{region.value}", signing_secret="s")
    return AgentServices(
        settings=SETTINGS,
        storage=StorageRouter(SETTINGS, overrides={region: storage}),
        scanner=NoopScanner(),
    )


async def _add_rate_card(
    database: Database, ctx: dict[str, Any], rows: list[tuple[str, RateUnit, str, str]]
) -> None:
    async with database.owner_session() as session:
        for category, unit, amount, currency in rows:
            session.add(
                RateCardEntry(
                    tenant_id=ctx["tenant_id"],
                    profile_id=ctx["profile_id"],
                    labor_category=category,
                    unit=unit,
                    rate_amount=Decimal(amount),
                    rate_currency=currency,
                )
            )


async def _run_pricing(
    database: Database, ctx: dict[str, Any], services: AgentServices
) -> tuple[uuid.UUID, Any]:
    specs, finish = pipeline.plan_steps("pricing")
    assert [s.agent for s in specs] == ["pricing"] and finish.status == "done"
    runner = AgentRunner(database, tenant_id=ctx["tenant_id"], llm=FakeLLM(), services=services)
    run_id = await runner.start(
        kind="pipeline", pursuit_id=ctx["pursuit_id"], params={"step": "pricing"}
    )
    return run_id, await runner.run(run_id, specs)


def _sheets(data: bytes) -> dict[str, list[list[Any]]]:
    book = openpyxl.load_workbook(io.BytesIO(data))
    try:
        return {
            name: [list(row) for row in book[name].iter_rows(values_only=True)]
            for name in book.sheetnames
        }
    finally:
        book.close()


async def test_us_pricing_template_uses_the_rate_card_and_placeholders_only(
    database: Database, tmp_path: Path
) -> None:
    services = _services(tmp_path)
    ctx = await _setup(database, requirements=US_REQUIREMENTS)
    await _add_rate_card(
        database,
        ctx,
        [
            ("Solutions Architect", RateUnit.HOUR, "210.50", "USD"),
            ("Cloud Engineer", RateUnit.HOUR, "165.00", "USD"),
        ],
    )
    run_id, result = await _run_pricing(database, ctx, services)
    assert result.status == "done", result.error
    out = PricingOutput.model_validate(result.outputs["pricing"])
    assert out.content_type == XLSX_CONTENT_TYPE and out.version == 1
    assert out.storage_key == pricing_template_key(ctx["tenant_id"], ctx["pursuit_id"], 1)
    assert out.file_name == "pricing-template-v1.xlsx"
    assert out.currency == "USD" and out.region == "us"
    assert out.rate_card_rows == 2 and out.labor_rows == 2 and out.warnings == []
    assert out.size_bytes > 0 and len(out.sha256) == 64

    data = await services.storage_for(Region.US).get(out.storage_key)
    assert len(data) == out.size_bytes
    sheets = _sheets(data)
    assert list(sheets) == ["Labor", "Placeholders", "Summary"]

    labor = sheets["Labor"]
    assert labor[0] == [
        "Labor category",
        "Unit",
        "Rate",
        "Currency",
        "Quantity",
        "Extended price",
    ]
    assert [row[0] for row in labor[1:]] == ["Cloud Engineer", "Solutions Architect"]
    assert [row[2] for row in labor[1:]] == [165.0, 210.5]  # straight from the rate card
    assert [row[4] for row in labor[1:]] == ["[NEEDS INPUT: hours]"] * 2
    assert [row[5] for row in labor[1:]] == ['=IFERROR(C2*E2,"")', '=IFERROR(C3*E3,"")']

    # SPEC 1 / SPEC 8: no price the rate card did not give. Every amount cell is a placeholder.
    holders = sheets["Placeholders"]
    assert holders[0] == ["Item", "Basis", "Amount", "Note"]
    assert {row[2] for row in holders[1:]} == {NEEDS_PRICE}
    assert len(holders) == 7  # header + six standard cost lines, no tax or fee lines
    summary = sheets["Summary"]
    assert [row[0] for row in summary[1:4]] == [
        "Labour subtotal",
        "Other direct costs",
        "Total",
    ]
    assert all(str(row[1]).startswith("=SUM(") for row in summary[1:4])
    # nothing anywhere in the workbook is a bare money number except the rate-card rates
    numbers = [
        cell
        for rows in sheets.values()
        for row in rows
        for cell in row
        if isinstance(cell, int | float)
    ]
    assert sorted(numbers) == [165.0, 210.5]

    async with database.session(ctx["tenant_id"]) as session:
        artifact = (await session.execute(select(PursuitArtifact))).scalar_one()
        assert artifact.kind == ARTIFACT_PRICING_TEMPLATE and artifact.version == 1
        assert artifact.created_by == "agent"
        assert artifact.data["storage_key"] == out.storage_key
        assert artifact.data["sha256"] == out.sha256
        step = (
            await session.execute(select(AgentStep).where(AgentStep.run_id == run_id))
        ).scalar_one()
        assert step.input_ref == f"pursuit:{ctx['pursuit_id']}:rate_card=2"
        assert step.tokens_in == 0 and step.cost_usd == 0  # no LLM: the step is free

    # a second run writes v2 next to v1 and never overwrites it
    _, again = await _run_pricing(database, ctx, services)
    second = PricingOutput.model_validate(again.outputs["pricing"])
    assert second.version == 2 and second.storage_key.endswith("/v2.xlsx")
    assert await services.storage_for(Region.US).exists(out.storage_key)


async def test_india_pricing_template_has_man_months_gst_and_emd_placeholders(
    database: Database, tmp_path: Path
) -> None:
    services = _services(tmp_path, Region.IN)
    ctx = await _setup(
        database,
        region=Region.IN,
        notice_type=NoticeType.GEM_BID,
        source_id="gem_bids",
        currency="INR",
        emd=Decimal("250000"),
    )
    await _add_rate_card(database, ctx, [("Project Manager", RateUnit.MONTH, "450000", "INR")])
    _, result = await _run_pricing(database, ctx, services)
    assert result.status == "done", result.error
    out = PricingOutput.model_validate(result.outputs["pricing"])
    assert out.currency == "INR" and out.region == "in"

    sheets = _sheets(await services.storage_for(Region.IN).get(out.storage_key))
    labor = sheets["Labor"]
    assert labor[1][:5] == ["Project Manager", "month", 450000, "INR", "[NEEDS INPUT: man-months]"]
    holders = {row[0]: row for row in sheets["Placeholders"][1:]}
    assert {"GST", "Earnest money deposit (EMD)", "Tender document fee"} <= set(holders)
    assert holders["GST"][2] == NEEDS_PRICE and "18%" in holders["GST"][3]
    assert holders["Earnest money deposit (EMD)"][2] == NEEDS_PRICE
    assert "₹2,50,000" in holders["Earnest money deposit (EMD)"][3]
    assert [row[0] for row in sheets["Summary"][1:6]] == [
        "Labour subtotal",
        "Other direct costs",
        "Bid costs (EMD, fees)",
        "Taxes",
        "Total",
    ]
    numbers = [
        cell
        for rows in sheets.values()
        for row in rows
        for cell in row
        if isinstance(cell, int | float)
    ]
    assert numbers == [450000]  # the one rate the tenant actually configured


async def test_pricing_without_a_rate_card_writes_placeholders_and_a_warning(
    database: Database, tmp_path: Path
) -> None:
    services = _services(tmp_path)
    ctx = await _setup(database)
    _, result = await _run_pricing(database, ctx, services)
    out = PricingOutput.model_validate(result.outputs["pricing"])
    assert out.rate_card_rows == 0 and out.warnings == [
        "the profile has no rate card, so every labour rate is a placeholder"
    ]
    sheets = _sheets(await services.storage_for(Region.US).get(out.storage_key))
    assert sheets["Labor"][1][0] == NEEDS_CATEGORY and sheets["Labor"][1][2] == NEEDS_PRICE
    assert not [
        cell
        for rows in sheets.values()
        for row in rows
        for cell in row
        if isinstance(cell, int | float)
    ]
    assert any("Warning: the profile has no rate card" in str(row[0]) for row in sheets["Summary"])


async def test_run_endpoint_queues_the_pipeline_and_pauses_at_the_first_missing_step(
    app: Any,
    api_client: httpx.AsyncClient,
    database: Database,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    app.state.agent_services = _services(tmp_path)
    ctx = await _setup(database)
    owner = auth_headers(user_id=ctx["user_id"], tenant_id=ctx["tenant_id"], email=ctx["email"])
    writer = auth_headers(user_id=uuid.uuid4(), tenant_id=ctx["tenant_id"], role=Role.WRITER)
    url = f"/api/v1/pursuits/{ctx['pursuit_id']}/agents/run"

    assert (await api_client.post(url, json={"step": "pricing"}, headers=writer)).status_code == 403
    bad = await api_client.post(url, json={"step": "nonsense"}, headers=owner)
    assert bad.status_code == 422

    monkeypatch.setattr("app.jobs.run_agents.enqueue_agents", lambda run_id, tenant_id: None)
    resp = await api_client.post(url, json={"step": "pricing", "inline": True}, headers=owner)
    assert resp.status_code == 202, resp.text
    body = resp.json()
    assert body["step"] == "pricing" and body["steps"] == ["pricing"] and body["mode"] == "inline"
    assert body["result"]["status"] == "done" and body["result"]["executed"] == ["pricing"]
    assert body["pursuit"]["run"]["id"] == body["run_id"]
    assert body["pursuit"]["run"]["status"] == "done"

    # "all" runs what exists and stops at the first agent a later task still has to add
    everything = await api_client.post(url, json={"step": "all", "inline": True}, headers=owner)
    assert everything.status_code == 202
    data = everything.json()
    assert data["steps"] == ["collect", "extract", "matrix"]
    assert data["result"]["status"] == "paused"
    assert "bid_no_bid" in data["result"]["pause_reason"]
    assert data["pursuit"]["run"]["status"] == "paused"

    # queued when a broker answers: the task id comes back and nothing runs in-process
    monkeypatch.setattr("app.jobs.run_agents.enqueue_agents", lambda run_id, tenant_id: "task-1")
    monkeypatch.setattr(SETTINGS, "celery_task_always_eager", False, raising=False)
    monkeypatch.setattr(app.state.settings, "celery_task_always_eager", False, raising=False)
    queued = await api_client.post(url, json={"step": "matrix"}, headers=owner)
    assert queued.status_code == 202
    assert queued.json()["mode"] == "queued" and queued.json()["task_id"] == "task-1"
    assert queued.json()["result"] is None
    assert queued.json()["pursuit"]["run"]["status"] == "queued"
