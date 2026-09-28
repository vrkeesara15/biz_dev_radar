"""M5-13: DOCX in the tenant's template, PDF from it, XLSX with the matrix / checklist /
pricing sheets and the ZIP named per the solicitation's rule -- every one carrying the
"DRAFT - internal" footer until the package is marked final.

The assertions open the real files with python-docx, pypdf and openpyxl (SPEC 12:
"exports open cleanly in Word / Acrobat / Excel").
"""

from __future__ import annotations

import dataclasses
import io
import uuid
import zipfile
from typing import Any

import httpx
import pytest
from app.core.citations import profile_token
from app.core.compliance import (
    ARTIFACT_CHECKLIST,
    ARTIFACT_FORMAT_RULES,
    ARTIFACT_OUTLINE,
    ARTIFACT_PRICING_TEMPLATE,
    ARTIFACT_RED_TEAM,
)
from app.core.config import Region
from app.core.db import Database
from app.core.disclaimers import VERIFY_ON_PORTAL
from app.core.exports import DRAFT_FOOTER
from app.core.opportunity import NoticeType
from app.core.profile_fields import PerformanceRole, ProfileFileKind
from app.core.roles import Role
from app.models import (
    AuditLog,
    CompanyProfile,
    ComplianceItem,
    Export,
    File,
    Opportunity,
    OpportunityDocument,
    PastPerformance,
    ProfileFile,
    Pursuit,
    PursuitArtifact,
    Requirement,
)
from app.services import exports as export_svc
from app.services.drafts import save_version
from app.services.storage import StorageRouter
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from tests.auth import auth_headers
from tests.factories import create_tenant_with_owner

SOLICITATION = "W91QUZ-26-R-0001"
NAMING_RULE = "{solicitation_number}_{volume}_{company}"

VOLUMES: list[dict[str, Any]] = [
    {
        "name": "Volume I - Technical",
        "sections": [
            {"id": "technical-approach", "title": "Technical Approach", "page_budget": 6},
            {"id": "quality-management", "title": "Quality Management", "page_budget": 2},
        ],
    },
    {
        "name": "Volume II - Past Performance",
        "sections": [{"id": "past-performance", "title": "Past Performance", "page_budget": 4}],
    },
]


def _pricing_workbook() -> bytes:
    workbook = Workbook()
    labor = workbook.active
    assert labor is not None
    labor.title = "Labor"
    labor.append(["Labor category", "Unit", "Rate", "Quantity"])
    labor.append(["Cloud engineer", "Hours", 185.0, "[NEEDS INPUT: hours]"])
    placeholders = workbook.create_sheet("Placeholders")
    placeholders.append(["[NEEDS INPUT: travel]", "[NEEDS INPUT: price]"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _setup(
    database: Database,
    settings: Any,
    *,
    with_template: bool = False,
    naming: str | None = NAMING_RULE,
    with_pricing: bool = True,
) -> dict[str, Any]:
    storage = StorageRouter(settings).for_region(Region.US)
    async with database.owner_session() as session:
        tenant, user, _ = await create_tenant_with_owner(session)
        profile = CompanyProfile(
            tenant_id=tenant.id, region=Region.US, legal_name="Alpha Federal LLC"
        )
        opp = Opportunity(
            source_id="sam_opps",
            external_id=f"ext-{uuid.uuid4().hex[:8]}",
            region=Region.US,
            country="US",
            currency="USD",
            notice_type=NoticeType.RFP,
            title="Cloud migration services",
            solicitation_number=SOLICITATION,
            source_url="https://sam.gov/opp/1",
        )
        session.add_all([profile, opp])
        await session.flush()
        past = PastPerformance(
            tenant_id=tenant.id,
            profile_id=profile.id,
            title="Treasury cloud migration",
            customer="US Treasury",
            role=PerformanceRole.PRIME,
            scope="Migrated 400 workloads",
        )
        doc = OpportunityDocument(
            opportunity_id=opp.id, url="https://x.test/rfp.pdf", file_name="rfp.pdf"
        )
        session.add_all([past, doc])
        await session.flush()
        pursuit = Pursuit(
            tenant_id=tenant.id,
            profile_id=profile.id,
            opportunity_id=opp.id,
            created_by=user.id,
            decision="bid",
            stage="drafting",
        )
        session.add(pursuit)
        await session.flush()
        req = Requirement(
            tenant_id=tenant.id,
            pursuit_id=pursuit.id,
            req_id="R-001",
            text="The contractor shall migrate 400 workloads.",
            document_id=doc.id,
            page=7,
            type="shall",
            quote="shall migrate 400 workloads",
        )
        session.add(req)
        await session.flush()
        session.add(
            ComplianceItem(
                tenant_id=tenant.id,
                pursuit_id=pursuit.id,
                requirement_id=req.id,
                section="Technical Approach",
                reason="keyword",
                status="drafted",
                notes="owner assigned",
            )
        )
        artifacts = [
            PursuitArtifact(
                tenant_id=tenant.id,
                pursuit_id=pursuit.id,
                kind=ARTIFACT_OUTLINE,
                version=1,
                data={
                    "outline": {
                        "volumes": VOLUMES,
                        "win_themes": [],
                        "unmapped_requirements": [],
                    }
                },
            ),
            PursuitArtifact(
                tenant_id=tenant.id,
                pursuit_id=pursuit.id,
                kind=ARTIFACT_FORMAT_RULES,
                version=1,
                data={"page_limit": 25, "font": "Times New Roman", "file_naming": naming},
            ),
            PursuitArtifact(
                tenant_id=tenant.id,
                pursuit_id=pursuit.id,
                kind=ARTIFACT_CHECKLIST,
                version=1,
                data={
                    "items": [
                        {
                            "key": "sam_registration",
                            "label": "Active SAM.gov registration",
                            "category": "registration",
                            "required": True,
                            "note": "check the CAGE code",
                            "source_req_ids": ["R-001"],
                        }
                    ]
                },
            ),
            PursuitArtifact(
                tenant_id=tenant.id,
                pursuit_id=pursuit.id,
                kind=ARTIFACT_RED_TEAM,
                version=1,
                data={"report": {"sections": [], "overall_score": 71, "missing_requirements": []}},
            ),
        ]
        if with_pricing:
            pricing_key = f"tenants/{tenant.id}/pursuits/{pursuit.id}/pricing/v1.xlsx"
            artifacts.append(
                PursuitArtifact(
                    tenant_id=tenant.id,
                    pursuit_id=pursuit.id,
                    kind=ARTIFACT_PRICING_TEMPLATE,
                    version=1,
                    data={"storage_key": pricing_key, "currency": "USD"},
                )
            )
        session.add_all(artifacts)
        await session.flush()
        if with_template:
            template = Document()
            template.add_paragraph("ALPHA FEDERAL PROPOSAL TEMPLATE")
            buffer = io.BytesIO()
            template.save(buffer)
            key = f"tenants/{tenant.id}/files/{uuid.uuid4()}.docx"
            await storage.put(
                key,
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            file_row = File(
                tenant_id=tenant.id,
                filename="template.docx",
                extension="docx",
                kind="docx",
                content_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
                size_bytes=len(buffer.getvalue()),
                sha256="0" * 64,
                region=Region.US,
                bucket=storage.bucket,
                key=key,
            )
            session.add(file_row)
            await session.flush()
            session.add(
                ProfileFile(
                    tenant_id=tenant.id,
                    profile_id=profile.id,
                    file_id=file_row.id,
                    kind=ProfileFileKind.TEMPLATE,
                    title="Proposal template",
                )
            )
            await session.flush()
        ctx = {
            "tenant_id": tenant.id,
            "user_id": user.id,
            "email": user.email,
            "profile_id": profile.id,
            "pursuit_id": pursuit.id,
            "past_performance": past.id,
        }
    if with_pricing:
        await storage.put(
            f"tenants/{ctx['tenant_id']}/pursuits/{ctx['pursuit_id']}/pricing/v1.xlsx",
            _pricing_workbook(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    token = profile_token("past_performance", ctx["past_performance"])
    ctx["token"] = token
    async with database.session(ctx["tenant_id"]) as session:
        await save_version(
            session,
            ctx["tenant_id"],
            ctx["pursuit_id"],
            "past-performance",
            title="Past Performance",
            volume="Volume II - Past Performance",
            body_html=f"<h2>Past Performance</h2><p>Treasury cloud migration [{token}].</p>",
            citations=[{"token": token, "quote": "Migrated 400 workloads"}],
            author="agent",
        )
        await save_version(
            session,
            ctx["tenant_id"],
            ctx["pursuit_id"],
            "technical-approach",
            title="Technical Approach",
            volume="Volume I - Technical",
            body_html=(
                f"<h2>Technical Approach</h2><p>We migrated 400 workloads [{token}].</p>"
                "<p>[NEEDS INPUT: name the transition manager]</p>"
            ),
            citations=[{"token": token, "quote": "Migrated 400 workloads"}],
            needs_input=[
                {"placeholder": "[NEEDS INPUT: name the transition manager]", "question": "Who?"}
            ],
            author="agent",
        )
    return ctx


async def _package(database: Database, settings: Any, ctx: dict[str, Any]) -> Any:
    storage = StorageRouter(settings).for_region(Region.US)
    async with database.session(ctx["tenant_id"]) as session:
        pursuit = await session.get(Pursuit, ctx["pursuit_id"])
        assert pursuit is not None
        return await export_svc.load_package(session, pursuit, storage=storage)


def _headers(ctx: dict[str, Any], role: Role | None = None) -> dict[str, str]:
    if role is None:
        return auth_headers(user_id=ctx["user_id"], tenant_id=ctx["tenant_id"], email=ctx["email"])
    return auth_headers(user_id=uuid.uuid4(), tenant_id=ctx["tenant_id"], role=role)


def _docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for section in document.sections:
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


# --- the package ------------------------------------------------------------------------


async def test_the_package_follows_the_outline_order(database: Database, settings: Any) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    assert [s.section_id for s in package.sections] == [
        "technical-approach",
        "past-performance",
    ], "the drafts were saved out of order; the outline decides"
    assert package.volumes == ["Volume I - Technical", "Volume II - Past Performance"]
    assert package.company == "Alpha Federal LLC"
    assert package.naming.template == NAMING_RULE
    assert [row.req_id for row in package.matrix] == ["R-001"]
    assert package.matrix[0].page == 7 and package.matrix[0].section == "Technical Approach"
    assert [item.key for item in package.checklist] == ["sam_registration"]
    assert package.pricing_xlsx is not None
    assert package.template_docx is None
    assert not package.final


# --- DOCX ---------------------------------------------------------------------------------


async def test_the_docx_opens_in_word_and_carries_the_draft_footer(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    data = export_svc.build_docx(package)

    document = Document(io.BytesIO(data))  # python-docx opening it IS the assertion
    text = _docx_text(data)
    assert "Cloud migration services" in text
    assert "Alpha Federal LLC" in text
    assert f"Solicitation {SOLICITATION}" in text
    # sections in outline order, each under its volume heading
    assert text.index("Volume I - Technical") < text.index("Technical Approach")
    assert text.index("Technical Approach") < text.index("Past Performance")
    assert "We migrated 400 workloads" in text
    assert f"References: [{ctx['token']}]" in text
    # SPEC 11: the internal footer and the portal disclaimer
    assert DRAFT_FOOTER in text
    assert VERIFY_ON_PORTAL in text
    assert "AI-generated draft" in text
    assert document.sections[0].footer.is_linked_to_previous is False

    # the [NEEDS INPUT] chip is highlighted yellow
    highlighted = [
        run.text
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.font.highlight_color is not None
    ]
    assert any("[NEEDS INPUT: name the transition manager]" in t for t in highlighted)


async def test_the_docx_uses_the_tenants_uploaded_template(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings, with_template=True)
    package = await _package(database, settings, ctx)
    assert package.template_docx is not None
    text = _docx_text(export_svc.build_docx(package))
    assert "ALPHA FEDERAL PROPOSAL TEMPLATE" in text
    assert "Technical Approach" in text


async def test_a_broken_template_falls_back_to_the_default(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    broken = dataclasses.replace(package, template_docx=b"not a docx")
    text = _docx_text(export_svc.build_docx(broken))
    assert "Technical Approach" in text and DRAFT_FOOTER in text


async def test_a_final_package_drops_the_draft_footer_but_keeps_the_disclaimer(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    final = dataclasses.replace(package, final=True)
    text = _docx_text(export_svc.build_docx(final))
    assert DRAFT_FOOTER not in text
    assert VERIFY_ON_PORTAL in text


# --- PDF ----------------------------------------------------------------------------------


async def test_the_pdf_opens_in_acrobat_and_carries_the_footer(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    rendered = export_svc.build_pdf(package, export_svc.build_docx(package))
    assert rendered.renderer in (export_svc.RENDERER_LIBREOFFICE, export_svc.RENDERER_PYMUPDF)

    reader = PdfReader(io.BytesIO(rendered.data))  # pypdf opening it IS the assertion
    assert len(reader.pages) >= 1
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Cloud migration services" in text
    assert "Technical Approach" in text
    assert VERIFY_ON_PORTAL.split(" on the")[0] in text
    assert "DRAFT - internal" in text


async def test_the_pure_python_fallback_is_used_when_libreoffice_is_missing(
    database: Database, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    monkeypatch.setattr(export_svc, "soffice_path", lambda: None)
    rendered = export_svc.build_pdf(package, export_svc.build_docx(package))
    assert rendered.renderer == export_svc.RENDERER_PYMUPDF
    reader = PdfReader(io.BytesIO(rendered.data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Alpha Federal LLC" in text and "DRAFT - internal" in text


# --- XLSX ---------------------------------------------------------------------------------


async def test_the_xlsx_opens_in_excel_with_matrix_checklist_and_pricing(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    workbook = load_workbook(io.BytesIO(export_svc.build_xlsx(package)))
    assert workbook.sheetnames == ["Matrix", "Checklist", "Pricing"]

    matrix = workbook["Matrix"]
    assert [c.value for c in matrix[1]][:5] == [
        "Requirement",
        "Type",
        "Page",
        "Text",
        "Section",
    ]
    assert matrix["A2"].value == "R-001" and matrix["C2"].value == 7
    assert matrix["E2"].value == "Technical Approach"

    checklist = workbook["Checklist"]
    assert checklist["A2"].value == "sam_registration"
    assert checklist["D2"].value == "yes" and checklist["F2"].value == "R-001"

    pricing = workbook["Pricing"]
    values = [row for row in pricing.iter_rows(values_only=True)]
    flat = [str(v) for row in values for v in row if v is not None]
    assert "Cloud engineer" in flat and "185" in flat
    assert any("[NEEDS INPUT" in v for v in flat)

    for name in workbook.sheetnames:
        sheet = workbook[name]
        last = str(sheet.cell(row=sheet.max_row, column=1).value or "")
        assert DRAFT_FOOTER in last and VERIFY_ON_PORTAL in last, name


async def test_the_xlsx_says_so_when_there_is_no_pricing_template(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings, with_pricing=False)
    package = await _package(database, settings, ctx)
    assert package.pricing_xlsx is None
    workbook = load_workbook(io.BytesIO(export_svc.build_xlsx(package)))
    assert "has not produced a template" in str(workbook["Pricing"]["A1"].value)


# --- ZIP ----------------------------------------------------------------------------------


async def test_the_zip_uses_the_solicitations_file_naming_rule(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    package = await _package(database, settings, ctx)
    rendered = export_svc.render(package, "zip")
    archive = zipfile.ZipFile(io.BytesIO(rendered.data))
    names = archive.namelist()
    assert names == [
        f"{SOLICITATION}_Technical_Alpha-Federal-LLC.docx",
        f"{SOLICITATION}_Technical_Alpha-Federal-LLC.pdf",
        f"{SOLICITATION}_Compliance_Alpha-Federal-LLC.xlsx",
        "MANIFEST.txt",
    ], names
    assert rendered.file_name == f"{SOLICITATION}_Package_Alpha-Federal-LLC.zip"
    manifest = archive.read("MANIFEST.txt").decode()
    assert "DRAFT - internal" in manifest and VERIFY_ON_PORTAL in manifest
    assert "Volume I - Technical / Technical Approach (v1, draft)" in manifest
    # the members are the real files
    Document(io.BytesIO(archive.read(names[0])))
    PdfReader(io.BytesIO(archive.read(names[1])))
    load_workbook(io.BytesIO(archive.read(names[2])))


async def test_a_notice_with_no_naming_rule_gets_the_default(
    database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings, naming=None)
    package = await _package(database, settings, ctx)
    archive = zipfile.ZipFile(io.BytesIO(export_svc.render(package, "zip").data))
    assert archive.namelist()[0] == f"{SOLICITATION}_Technical_Alpha-Federal-LLC.docx"


# --- the routes ------------------------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["docx", "pdf", "xlsx", "zip"])
async def test_the_export_route_stores_a_row_and_a_signed_url(
    api_client: httpx.AsyncClient, database: Database, settings: Any, fmt: str
) -> None:
    ctx = await _setup(database, settings)
    answer = await api_client.post(
        f"/api/v1/pursuits/{ctx['pursuit_id']}/export?format={fmt}", headers=_headers(ctx)
    )
    assert answer.status_code == 202, answer.text
    body = answer.json()
    assert body["format"] == fmt and body["version"] == 1
    assert body["size_bytes"] > 0 and body["final"] is False
    assert body["file_name"].endswith(f".{fmt}")
    assert body["url"] and body["expires_in"] > 0
    assert (body["renderer"] is None) == (fmt in ("docx", "xlsx"))

    async with database.owner_session() as session:
        rows = list((await session.execute(select(Export))).scalars().all())
        assert len(rows) == 1
        assert rows[0].storage_key.endswith(f".{fmt}")
        assert rows[0].created_by == ctx["user_id"]
        audits = list(
            (await session.execute(select(AuditLog).where(AuditLog.action == "export.created")))
            .scalars()
            .all()
        )
    assert len(audits) == 1
    assert audits[0].meta["format"] == fmt and audits[0].meta["sections"] == 2

    # the stored object really is the file
    storage = StorageRouter(settings).for_region(Region.US)
    data = await storage.get(rows[0].storage_key)
    assert len(data) == body["size_bytes"]


async def test_exports_are_versioned_listed_and_the_download_is_audited(
    api_client: httpx.AsyncClient, database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    base = f"/api/v1/pursuits/{ctx['pursuit_id']}"
    first = await api_client.post(f"{base}/export?format=docx", headers=_headers(ctx))
    second = await api_client.post(f"{base}/export?format=docx", headers=_headers(ctx))
    assert [first.json()["version"], second.json()["version"]] == [1, 2]

    listed = await api_client.get(f"{base}/exports", headers=_headers(ctx))
    assert listed.status_code == 200
    assert listed.json()["count"] == 2 and listed.json()["package_final"] is False
    assert all(item["url"] is None for item in listed.json()["items"]), "listing hands out no URL"

    export_id = second.json()["id"]
    download = await api_client.get(f"{base}/exports/{export_id}", headers=_headers(ctx))
    assert download.status_code == 200
    assert download.json()["url"] and download.json()["expires_at"]
    missing = await api_client.get(f"{base}/exports/{uuid.uuid4()}", headers=_headers(ctx))
    assert missing.status_code == 404

    async with database.owner_session() as session:
        reads = list(
            (await session.execute(select(AuditLog).where(AuditLog.action == "export.read")))
            .scalars()
            .all()
        )
    assert len(reads) == 1 and reads[0].meta["version"] == 2
    assert reads[0].user_id == ctx["user_id"]


async def test_export_roles_and_a_bad_format(
    api_client: httpx.AsyncClient, database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    base = f"/api/v1/pursuits/{ctx['pursuit_id']}"
    viewer = _headers(ctx, Role.VIEWER)
    assert (await api_client.post(f"{base}/export?format=docx", headers=viewer)).status_code == 403
    assert (await api_client.get(f"{base}/exports", headers=viewer)).status_code == 403
    reviewer = await api_client.post(
        f"{base}/export?format=docx", headers=_headers(ctx, Role.REVIEWER)
    )
    assert reviewer.status_code == 202
    bad = await api_client.post(f"{base}/export?format=pptx", headers=_headers(ctx))
    assert bad.status_code == 422


async def test_mark_final_needs_gate_two_and_then_drops_the_draft_footer(
    api_client: httpx.AsyncClient, database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    base = f"/api/v1/pursuits/{ctx['pursuit_id']}"

    early = await api_client.post(f"{base}/mark-final", json={}, headers=_headers(ctx))
    assert early.status_code == 409 and "Gate 2" in early.json()["detail"]

    approved = await api_client.post(
        f"{base}/approve-package", json={"inline": True}, headers=_headers(ctx)
    )
    assert approved.status_code == 200, approved.text

    assert (
        await api_client.post(f"{base}/mark-final", json={}, headers=_headers(ctx, Role.WRITER))
    ).status_code == 403

    final = await api_client.post(
        f"{base}/mark-final", json={"note": "signed off"}, headers=_headers(ctx)
    )
    assert final.status_code == 200, final.text
    assert final.json()["package_final"] is True
    assert final.json()["package_final_by"] == str(ctx["user_id"])

    exported = await api_client.post(f"{base}/export?format=docx", headers=_headers(ctx))
    assert exported.json()["final"] is True
    storage = StorageRouter(settings).for_region(Region.US)
    async with database.owner_session() as session:
        row = (
            await session.execute(
                select(Export).where(Export.id == uuid.UUID(exported.json()["id"]))
            )
        ).scalar_one()
    text = _docx_text(await storage.get(row.storage_key))
    assert DRAFT_FOOTER not in text
    assert VERIFY_ON_PORTAL in text

    async with database.owner_session() as session:
        audits = list(
            (
                await session.execute(
                    select(AuditLog).where(AuditLog.action == "pursuit.package_final")
                )
            )
            .scalars()
            .all()
        )
    assert len(audits) == 1 and audits[0].meta["note"] == "signed off"


async def test_another_tenant_cannot_export_or_download(
    api_client: httpx.AsyncClient, database: Database, settings: Any
) -> None:
    ctx = await _setup(database, settings)
    created = await api_client.post(
        f"/api/v1/pursuits/{ctx['pursuit_id']}/export?format=docx", headers=_headers(ctx)
    )
    export_id = created.json()["id"]
    async with database.owner_session() as session:
        other, other_user, _ = await create_tenant_with_owner(session)
    foreign = auth_headers(user_id=other_user.id, tenant_id=other.id, email=other_user.email)
    base = f"/api/v1/pursuits/{ctx['pursuit_id']}"
    assert (await api_client.post(f"{base}/export?format=docx", headers=foreign)).status_code == 404
    assert (await api_client.get(f"{base}/exports", headers=foreign)).status_code == 404
    assert (await api_client.get(f"{base}/exports/{export_id}", headers=foreign)).status_code == 404
    assert (
        await api_client.post(f"{base}/mark-final", json={}, headers=foreign)
    ).status_code == 404
