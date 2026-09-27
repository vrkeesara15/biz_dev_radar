"""Agent 7: pricing template builder (SPEC 8, 1). Pipeline step "pricing", no LLM.

Renders app.core.pricing's plan as an XLSX with three sheets -- Labor (the tenant's rate
card, one row per labour category, with the quantity left blank), Placeholders (every
price the rate card does not supply, India's EMD, tender fee, bank-guarantee charges and
GST among them) and Summary (SUM formulas over both) -- stores it through the region's
Storage and records it as a `pricing_template` pursuit artifact.

SPEC 1 puts cost-volume generation out of scope: this agent never invents a price. Rate
cells come from rate_card_entries and nothing else; every other money cell is a yellow
[NEEDS INPUT: price].
"""

from __future__ import annotations

import hashlib
import io
import uuid
from decimal import Decimal
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field
from sqlalchemy import select

from app.agents.pipeline import STEP_PRICING, register
from app.agents.runner import StepContext
from app.core.compliance import ARTIFACT_PRICING_TEMPLATE
from app.core.pricing import (
    LABOR_HEADERS,
    PLACEHOLDER_HEADERS,
    SHEET_LABOR,
    SHEET_PLACEHOLDERS,
    SHEET_SUMMARY,
    SUMMARY_HEADERS,
    PricingPlan,
    RateCardRow,
    build_plan,
)
from app.models import Opportunity, Pursuit, RateCardEntry
from app.services.pursuits import store_artifact

log = structlog.get_logger(__name__)

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PLACEHOLDER_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow: a human must fill it
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
MONEY_FORMAT = "#,##0.00"
COLUMN_WIDTHS = (38, 12, 14, 10, 16, 18)
DRAFT_NOTE = (
    "Template only. Every yellow cell needs a price or quantity from a human; BidRadar "
    "never fills a price that is not on the tenant's rate card."
)


class PricingOutput(BaseModel):
    storage_key: str
    file_name: str
    content_type: str = XLSX_CONTENT_TYPE
    sha256: str
    size_bytes: int
    currency: str
    region: str
    labor_rows: int = 0
    rate_card_rows: int = 0
    placeholder_cells: int = 0
    version: int | None = None
    warnings: list[str] = Field(default_factory=list)


def pricing_template_key(tenant_id: uuid.UUID, pursuit_id: uuid.UUID, version: int) -> str:
    """tenants/{tenant}/pursuits/{pursuit}/pricing/v{n}.xlsx (the bucket carries residency)."""
    return f"tenants/{tenant_id}/pursuits/{pursuit_id}/pricing/v{version}.xlsx"


def _header(sheet: Worksheet, headers: tuple[str, ...]) -> None:
    sheet.append(list(headers))
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    sheet.freeze_panes = "A2"


def _widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _placeholder(sheet: Worksheet, row: int, column: int, text: str) -> None:
    cell = sheet.cell(row=row, column=column, value=text)
    cell.fill = PLACEHOLDER_FILL
    cell.alignment = Alignment(horizontal="left")


def render_workbook(plan: PricingPlan, *, title: str | None = None) -> bytes:
    """The XLSX bytes for a plan. Pure apart from openpyxl's in-memory serialisation."""
    book = Workbook()
    labor = book.active
    assert labor is not None
    labor.title = SHEET_LABOR
    _header(labor, LABOR_HEADERS)
    _widths(labor, COLUMN_WIDTHS)
    for index, row in enumerate(plan.labor, start=2):
        labor.cell(row=index, column=1, value=row.category)
        labor.cell(row=index, column=2, value=row.unit)
        if row.rate is None:
            _placeholder(labor, index, 3, row.placeholder or "")
        else:
            rate = labor.cell(row=index, column=3, value=float(row.rate))
            rate.number_format = MONEY_FORMAT
        labor.cell(row=index, column=4, value=row.currency or plan.currency)
        _placeholder(labor, index, 5, row.quantity_placeholder)
        extended = labor.cell(row=index, column=6, value=f'=IFERROR(C{index}*E{index},"")')
        extended.number_format = MONEY_FORMAT

    holders = book.create_sheet(SHEET_PLACEHOLDERS)
    _header(holders, PLACEHOLDER_HEADERS)
    _widths(holders, (38, 24, 18, 60))
    for index, item in enumerate(plan.placeholders, start=2):
        holders.cell(row=index, column=1, value=item.label)
        holders.cell(row=index, column=2, value=item.basis)
        _placeholder(holders, index, 3, item.amount)
        holders.cell(row=index, column=4, value=item.note)

    summary = book.create_sheet(SHEET_SUMMARY)
    _header(summary, SUMMARY_HEADERS)
    _widths(summary, (38, 20))
    for index, line in enumerate(plan.summary, start=2):
        summary.cell(row=index, column=1, value=line.label)
        cell = summary.cell(row=index, column=2, value=line.formula)
        cell.number_format = MONEY_FORMAT
    note_row = len(plan.summary) + 3
    summary.cell(row=note_row, column=1, value=DRAFT_NOTE)
    summary.cell(row=note_row + 1, column=1, value=f"Currency: {plan.currency}")
    if title:
        summary.cell(row=note_row + 2, column=1, value=title)
    for offset, warning in enumerate(plan.warnings):
        summary.cell(row=note_row + 3 + offset, column=1, value=f"Warning: {warning}")

    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


# --- pipeline step -------------------------------------------------------------------


async def _pursuit_rate_card(
    session: Any, pursuit_id: uuid.UUID | None
) -> tuple[Pursuit, Opportunity, list[RateCardRow]]:
    if pursuit_id is None:
        raise RuntimeError("pricing needs a pursuit")
    pursuit = await session.get(Pursuit, pursuit_id)
    if pursuit is None:
        raise LookupError(f"pursuit {pursuit_id} not found")
    opportunity = await session.get(Opportunity, pursuit.opportunity_id)
    if opportunity is None:
        raise LookupError(f"opportunity {pursuit.opportunity_id} not found")
    entries = (
        (
            await session.execute(
                select(RateCardEntry)
                .where(RateCardEntry.profile_id == pursuit.profile_id)
                .order_by(RateCardEntry.labor_category)
            )
        )
        .scalars()
        .all()
    )
    rate_card = [
        RateCardRow(
            labor_category=entry.labor_category,
            unit=str(entry.unit),
            rate_amount=Decimal(entry.rate_amount),
            rate_currency=entry.rate_currency,
            min_years_experience=entry.min_years_experience,
            notes=entry.notes,
        )
        for entry in entries
    ]
    return pursuit, opportunity, rate_card


@register(STEP_PRICING)
async def pricing(ctx: StepContext) -> PricingOutput:
    services = ctx.require_services()
    pursuit, opportunity, rate_card = await _pursuit_rate_card(ctx.session, ctx.run.pursuit_id)
    ctx.step.input_ref = f"pursuit:{pursuit.id}:rate_card={len(rate_card)}"
    plan = build_plan(
        str(opportunity.region),
        rate_card,
        emd_amount=opportunity.emd_amount,
        tender_fee=opportunity.tender_fee,
    )
    data = render_workbook(plan, title=opportunity.title)
    artifact = await store_artifact(
        ctx.session,
        ctx.tenant_id,
        pursuit.id,
        ARTIFACT_PRICING_TEMPLATE,
        {"pending": True},
    )
    key = pricing_template_key(ctx.tenant_id, pursuit.id, artifact.version)
    await services.storage_for(opportunity.region).put(key, data, XLSX_CONTENT_TYPE)
    output = PricingOutput(
        storage_key=key,
        file_name=f"pricing-template-v{artifact.version}.xlsx",
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        currency=plan.currency,
        region=plan.region,
        labor_rows=len(plan.labor),
        rate_card_rows=len(rate_card),
        placeholder_cells=len(plan.amounts()) + len(plan.labor),  # + one quantity per row
        version=artifact.version,
        warnings=list(plan.warnings),
    )
    artifact.data = output.model_dump(mode="json")
    await ctx.session.flush()
    log.info(
        "pricing.done",
        pursuit_id=str(pursuit.id),
        version=artifact.version,
        rate_card_rows=len(rate_card),
        placeholders=output.placeholder_cells,
    )
    return output
