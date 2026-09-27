"""XLSX -> one page per worksheet (rows tab-separated), one section per sheet."""

from __future__ import annotations

import hashlib
import io

from app.core.parsing.types import Page, ParsedDocument, ParseError, Section


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def parse_xlsx(data: bytes, *, max_rows_per_sheet: int = 5000) -> ParsedDocument:
    import openpyxl

    doc = ParsedDocument(kind="xlsx", sha256=hashlib.sha256(data).hexdigest(), parser="openpyxl")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"openpyxl: {exc}") from exc
    try:
        for index, sheet in enumerate(workbook.worksheets, start=1):
            lines: list[str] = []
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_no > max_rows_per_sheet:
                    doc.warnings.append(f"{sheet.title}: truncated at {max_rows_per_sheet} rows")
                    break
                cells = [_cell(v) for v in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip("\t"))
            text = "\n".join(lines)
            doc.pages.append(Page(number=index, text=text))
            doc.sections.append(Section(sheet.title, index, index, text))
    finally:
        workbook.close()
    doc.page_count = len(doc.pages)
    return doc
